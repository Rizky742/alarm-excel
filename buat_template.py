#!/usr/bin/env python3
"""
Buat template Excel untuk import ke Free Alarm Clock.
Jalankan: python3 buat_template.py
Hasil: contoh_jadwal.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Jadwal"

headers = ["Waktu (HH:MM)", "Label", "Hari", "Tipe", "Suara", "Pesan"]
widths = [16, 30, 40, 12, 25, 30]
for col, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="4472C4")
    c.alignment = Alignment(horizontal="center")
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

contoh = [
    ["07:00", "Bel Masuk Sekolah", "Senin,Selasa,Rabu,Kamis,Jumat", "Weekly", "School.mp3", "Selamat pagi, apel pagi"],
    ["09:30", "Istirahat", "Senin,Selasa,Rabu,Kamis,Jumat", "Weekly", "Bells.mp3", "Waktunya istirahat"],
    ["12:00", "Pulang", "Senin,Kamis", "Weekly", "School.mp3", "Pulang sekolah"],
    ["07:30", "Upacara Bendera", "Senin", "Weekly", "Classic.mp3", "Upacara hari Senin"],
    ["08:00", "Rapat Sekali Saja", "2026-10-05", "Once", "Alarm.mp3", "Rapat penting"],
    ["06:30", "Bangun Pagi Harian", "Setiap Hari", "Daily", "Rooster.mp3", "Bangun!"],
]

for r, row in enumerate(contoh, 2):
    for col, val in enumerate(row, 1):
        ws.cell(row=r, column=col, value=val)

ws2 = wb.create_sheet("Petunjuk")
petunjuk = [
    ["Kolom", "Isi yang benar"],
    ["Waktu", "Format HH:MM 24-jam, contoh 07:00, 13:30"],
    ["Label", "Nama alarm, contoh: Bel Masuk"],
    ["Hari", "Untuk Tipe=Weekly: Senin,Selasa,Rabu,Kamis,Jumat,Sabtu,Minggu / Weekdays / Weekends / Setiap Hari"],
    ["", "Untuk Tipe=Once: tanggal YYYY-MM-DD, contoh 2026-10-05"],
    ["", "Untuk Tipe=Daily: kosongkan atau 'Setiap Hari'"],
    ["Tipe", "Once / Daily / Weekly (Monthly/Yearly juga didukung FreeAlarmClock v5.3)"],
    ["Suara", "Nama file di folder Sounds\\ misal School.mp3, atau path MP3 lengkap"],
    ["Pesan", "Opsional, ditampilkan saat alarm bunyi"],
]
for r, row in enumerate(petunjuk, 1):
    for c, v in enumerate(row, 1):
        ws2.cell(row=r, column=c, value=v)
ws2.column_dimensions["A"].width = 12
ws2.column_dimensions["B"].width = 90

wb.save("/Users/macbook2/freealarm-excel-import/contoh_jadwal.xlsx")
print("OK: contoh_jadwal.xlsx dibuat")
